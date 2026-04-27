from docx import Document


def main() -> None:
    doc = Document()

    doc.add_heading("Hairvelous User Survey", level=1)
    doc.add_paragraph(
        "Source: "
        "https://docs.google.com/forms/d/e/1FAIpQLSfBBxJ6tvR2U5I28DBm0gFN0_TZz926aHJDRh_ZSlme6-7LpQ/viewform?usp=header"
    )

    doc.add_paragraph("")
    doc.add_paragraph(
        'Thank you for participating in this survey. We are conducting a study for our '
        'capstone project titled "Hairvelous: Hair Care Consultation and Product '
        'Recommendation System."'
    )
    doc.add_paragraph(
        "Your responses will help us understand user needs, hair concerns, and feature "
        "preferences. All information will remain confidential and used only for academic "
        "purposes."
    )

    doc.add_paragraph("")
    doc.add_heading("Instructions", level=2)
    doc.add_paragraph("Please answer all questions honestly.", style="List Bullet")
    doc.add_paragraph("Your responses are confidential.", style="List Bullet")
    doc.add_paragraph("This survey is for academic purposes only.", style="List Bullet")

    doc.add_paragraph("")
    doc.add_heading("Notes", level=2)
    doc.add_paragraph("Sign in to Google to save your progress.", style="List Bullet")
    doc.add_paragraph("Never submit passwords through Google Forms.", style="List Bullet")
    doc.add_paragraph(
        "This content is neither created nor endorsed by Google.",
        style="List Bullet",
    )

    doc.save("Hairvelous_User_Survey.docx")


if __name__ == "__main__":
    main()
